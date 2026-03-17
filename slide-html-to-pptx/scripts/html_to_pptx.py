#!/usr/bin/env python3
"""
slide-html-to-pptx skill — Convert HTML slide decks to editable PowerPoint.

Usage:
    python skills/html_to_pptx.py --input-dir ./slides --output deck.pptx
    python skills/html_to_pptx.py --slides slide-01.html,slide-02.html --output deck.pptx
    python skills/html_to_pptx.py --manifest slides.json --output deck.pptx

The manifest JSON format:
    [
        {"file": "slide-01.html", "title": "Title Slide", "act": "Intro"},
        {"file": "slide-02.html", "title": "Hook Statement", "act": "Intro"},
        ...
    ]
"""

import argparse, os, sys, json, tempfile, glob
from pathlib import Path

# Import from sibling skills
_root = Path(__file__).parent.parent.parent
sys.path.insert(0, str(_root / "chrome-extract" / "scripts"))
sys.path.insert(0, str(_root / "slide-pptx-builder" / "scripts"))
sys.path.insert(0, str(_root / "slide-html-standardize" / "scripts"))

from chrome_extract import extract_layout, screenshot_html, crop_region
from pptx_builder import SlideBuilder
from html_standardize import standardize_html
from pptx import Presentation
from pptx.util import Inches
from bs4 import BeautifulSoup


def extract_notes(html_path):
    """Extract text from HTML for speaker notes."""
    try:
        with open(html_path, 'r', encoding='utf-8') as f:
            soup = BeautifulSoup(f.read(), 'lxml')
        for tag in soup.find_all(['script', 'style', 'nav']):
            tag.decompose()
        sl = soup.find(class_='slide') or soup.body
        return sl.get_text(separator='\n', strip=True)[:3000] if sl else ''
    except:
        return ''


def build_deck(slides, input_dir, output_path, hide_selectors=None,
               viewport_w=1280, viewport_h=720):
    """
    Build a PowerPoint deck from HTML slides.

    Args:
        slides: list of dicts with at least 'file' key, optionally 'title', 'act'
        input_dir: directory containing HTML files
        output_path: output .pptx path
        hide_selectors: CSS selectors to hide during extraction
        viewport_w/h: browser viewport size
    """
    sw, sh = 13.333, 7.5
    builder = SlideBuilder(sw, sh, viewport_w, viewport_h)
    prs = Presentation()
    prs.slide_width = Inches(sw)
    prs.slide_height = Inches(sh)

    total = len(slides)
    stats = {'layout': 0, 'fallback': 0}

    with tempfile.TemporaryDirectory() as tmpdir:
        for i, si in enumerate(slides):
            filename = si['file']
            title = si.get('title', filename)
            act = si.get('act', '')
            html_path = os.path.join(input_dir, filename)

            if not os.path.exists(html_path):
                print(f"  [{i+1}/{total}] {filename}: NOT FOUND, skipping")
                continue

            print(f"  [{i+1}/{total}] {filename}: {title}", end="", flush=True)

            try:
                layout = extract_layout(html_path, tmpdir, viewport_w, viewport_h,
                                       hide_selectors)

                if layout and 'error' not in layout and len(layout.get('elements', [])) > 0:
                    # Create SVG screenshot callback
                    full_ss_path = [None]  # mutable for closure
                    def screenshot_svg(svg_rect, _si=si, _fss=full_ss_path):
                        if not _fss[0]:
                            _fss[0] = os.path.join(tmpdir, f"full_{_si['file']}.png")
                            screenshot_html(os.path.join(input_dir, _si['file']),
                                          _fss[0], tmpdir, viewport_w, viewport_h,
                                          hide_selectors=hide_selectors)
                        out = os.path.join(tmpdir,
                            f"svg_{i}_{int(svg_rect['x'])}_{int(svg_rect['y'])}.png")
                        if crop_region(_fss[0], out, svg_rect, viewport_w, viewport_h):
                            return out
                        return None

                    notes = f"[{act}] {title}\n\n{extract_notes(html_path)}" if act else extract_notes(html_path)
                    builder.build_slide(prs, layout, screenshot_svg, notes)

                    ns = len([e for e in layout['elements'] if e['type'] == 'shape'])
                    nt = len([e for e in layout['elements'] if e['type'] in ('text', 'richtext')])
                    nv = len(layout.get('svgElements', []))
                    print(f" [{ns}s {nt}t {nv}svg]")
                    stats['layout'] += 1
                else:
                    # Fallback: screenshot-based slide
                    _build_fallback(prs, builder, html_path, si, tmpdir,
                                   viewport_w, viewport_h, hide_selectors)
                    print(" [fallback]")
                    stats['fallback'] += 1
            except Exception as e:
                try:
                    _build_fallback(prs, builder, html_path, si, tmpdir,
                                   viewport_w, viewport_h, hide_selectors)
                    print(f" [fallback: {e}]")
                except:
                    print(f" [ERROR: {e}]")
                stats['fallback'] += 1

    prs.save(output_path)
    print(f"\nSaved: {output_path}")
    print(f"Layout: {stats['layout']}/{total}, Fallback: {stats['fallback']}")
    return output_path


def _build_fallback(prs, builder, html_path, si, tmpdir, vw, vh, hide_selectors):
    """Screenshot fallback when extraction fails."""
    from pptx.dml.color import RGBColor
    from pptx.util import Emu
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0x0A, 0x0A, 0x0F)
    png = os.path.join(tmpdir, f"fb_{os.path.basename(html_path)}.png")
    if screenshot_html(html_path, png, tmpdir, vw, vh, hide_selectors=hide_selectors):
        slide.shapes.add_picture(png, Emu(0), Emu(0),
                                Inches(builder.sw), Inches(builder.sh))
    notes = extract_notes(html_path)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def main():
    parser = argparse.ArgumentParser(description='Convert HTML slides to editable PowerPoint')
    parser.add_argument('--input-dir', '-d', default='.', help='Directory containing HTML files')
    parser.add_argument('--output', '-o', default='deck.pptx', help='Output PPTX path')
    parser.add_argument('--slides', '-s', help='Comma-separated HTML filenames')
    parser.add_argument('--manifest', '-m', help='JSON manifest file with slide list')
    parser.add_argument('--hide', help='CSS selectors to hide (comma-separated)')
    parser.add_argument('--viewport', default='1280x720', help='Viewport WxH (default 1280x720)')
    parser.add_argument('--standardize', action='store_true',
                        help='Run slide-html-standardize on slides before conversion')
    args = parser.parse_args()

    vw, vh = map(int, args.viewport.split('x'))
    hide = args.hide.split(',') if args.hide else ['.slide-nav']

    if args.manifest:
        with open(args.manifest) as f:
            slides = json.load(f)
    elif args.slides:
        slides = [{'file': f.strip()} for f in args.slides.split(',')]
    else:
        # Auto-discover HTML files
        html_files = sorted(glob.glob(os.path.join(args.input_dir, '*.html')))
        slides = [{'file': os.path.basename(f)} for f in html_files
                  if 'index' not in os.path.basename(f).lower()
                  and 'sorter' not in os.path.basename(f).lower()]

    if not slides:
        print("No slides found. Use --slides, --manifest, or place HTML files in --input-dir")
        sys.exit(1)

    if args.standardize:
        print(f"Standardizing {len(slides)} slides...")
        for si in slides:
            html_path = os.path.join(args.input_dir, si['file'])
            if os.path.exists(html_path):
                standardize_html(html_path, viewport_w=vw, viewport_h=vh)

    print(f"Converting {len(slides)} slides from {args.input_dir}")
    build_deck(slides, args.input_dir, args.output, hide, vw, vh)


if __name__ == '__main__':
    main()
